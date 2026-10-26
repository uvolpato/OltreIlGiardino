"""Converte il file articoli nei dati del prototipo del catalogo.
Uso:  python Catalogo/converti.py

Legge: Documenti Vari/Articoli IBED_13_10_25_rev2_mod.xlsx
Scrive: Prototipo-Catalogo/assets/articoli.js  (window.CAT = {articoli, colori, classi, materiali, anomalie})
"""
import json, re, collections
from pathlib import Path
import openpyxl

RADICE = Path(__file__).resolve().parent.parent
SRC = next((RADICE / "Documenti Vari").glob("Articoli*13_10_25_rev2_mod.xlsx"))   # oggi si chiama "Articoli IBED_13_10_25_rev2_mod.xlsx"
DEST = RADICE / "Prototipo-Catalogo" / "assets" / "articoli.js"
SCHEMA = re.compile(r"^([A-Z]{2})([A-Z]{3})(\d{2})(\d{3})-?(\d{2})?$")
# "/N" dopo una misura = massimo (h45/72: altezza regolabile da 45 a 72); non lo è se apre un'altra misura ("…h45/80x80xh60")
N = r"\d+(?:[.,]\d+)?"
MX = lambda g, fine=True: rf"(?:\s*/\s*(?P<{g}>{N})(?![\d.,]){r'(?!\s*[xX×])' if fine else ''})?"   # fine: l'ultima misura non può aprirne un'altra; (?![\d.,]) evita di accorciare "80" in "8"
DIM = re.compile(rf"(?:(?P<d>[ØøΦ⌀]|diam\.?\s*)\s*(?P<dv>{N}){MX('dx')}|(?P<l>{N}){MX('lx', False)}\s*[xX×]\s*(?P<p>{N}){MX('px', False)})(?:\s*[xX×]?\s*h\s*(?P<h>{N}){MX('hx')})?", re.I)

wb = openpyxl.load_workbook(SRC, data_only=True)

# tabella colori: codice -> nome (dalle voci "01 Bianco" ecc.) e famiglia
colori, materiali = {}, {}
blocco = "colori"
for r in wb["Colori"].iter_rows(values_only=True):
    v = [c for c in r if c is not None]
    if not v:
        continue
    if v[0] == "Descrizione" and colori:
        blocco = "materiali"
        continue
    if v[0] == "Descrizione":
        continue
    fam, voci = str(v[0]).strip(), v[3:]
    dest = colori if blocco == "colori" else materiali
    dest.setdefault("_fam", {})[str(v[1]).strip()] = fam
    for x in voci:
        m = re.match(r"\s*(\d{2})\s*(.+)", str(x))
        if m:
            dest[m.group(1)] = m.group(2).strip().replace("�", "ù")
colfam = colori.pop("_fam"); matfam = materiali.pop("_fam")

# classificazioni con codici
classi = []
for r in wb["Classificazioni"].iter_rows(min_row=3, values_only=True):
    if r[4] and r[6] and r[8]:
        classi.append([r[4], r[5], r[6], r[7], r[8], r[9]])

def num(x):
    return float(x.replace(",", ".")) if x else None

arts, anom = [], collections.defaultdict(list)
righe = [r for r in wb["ARTICOLI"].iter_rows(min_row=3, values_only=True) if r[0]]
cont = collections.Counter(re.sub(r"\s", "", str(r[0])) for r in righe)
for r in righe:
    raw = str(r[0]); c = re.sub(r"\s", "", raw); m = SCHEMA.match(c)
    base = c.rsplit("-", 1)[0] if "-" in c else c
    a = {"c": c, "b": base, "d": (r[1] or "").strip(), "t": (r[3] or "").strip(), "e": (r[4] or "").strip(),
         "dim": re.sub(r"\s+", " ", str(r[5] or "")).strip(), "f": (r[7] or "").strip(), "f2": (r[9] or "").strip(),
         "n": (r[10] or "").strip() if isinstance(r[10], str) else "", "um": r[11] or "NR",
         "pn": r[12] if isinstance(r[12], (int, float)) else None, "pv": r[14] if isinstance(r[14], (int, float)) else None,
         "k1": (r[16] or "").strip(), "k2": (r[17] or "").strip(), "k3": (r[18] or "").strip(),
         "foto": [x for x in r[19:24] if x]}
    if m:
        a["mat"] = m.group(3); a["cc"] = m.group(5) or ""
    else:
        anom["fuori_schema"].append(c)
    d = DIM.search(a["dim"])
    if d:
        if d.group("dv"): a["D"] = num(d.group("dv"))
        else: a["L"], a["P"] = num(d.group("l")), num(d.group("p"))
        if d.group("h"): a["H"] = num(d.group("h"))
        for k, g in (("Dx", "dx"), ("Lx", "lx"), ("Px", "px"), ("Hx", "hx")):   # massimo solo se maggiore del minimo
            if d.group(g) and num(d.group(g)) > (a.get(k[0]) or 0): a[k] = num(d.group(g))
    elif a["dim"]:
        anom["dimensioni"].append(c)
    if cont[c] > 1:
        a["dup"] = True; anom["duplicati"].append(c)
    if len(a["d"]) > 40:
        anom["breve_lunga"].append(c)
    if re.search(r"copertur|fodera|rivestiment", a["d"], re.I):
        anom["coperture"].append(c)
    arts.append({k: v for k, v in a.items() if v not in ("", None, [])})

anom = {k: sorted(set(v)) for k, v in anom.items()}
DEST.parent.mkdir(parents=True, exist_ok=True)
DEST.write_text("/* generato da Catalogo/converti.py */\nwindow.CAT=" + json.dumps(
    {"articoli": arts, "colori": colori, "colfam": colfam, "materiali": materiali, "matfam": matfam, "classi": classi, "anomalie": anom},
    ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
print(f"{len(arts)} articoli, {len(colori)} colori, {len(classi)} classi ->", DEST, f"{DEST.stat().st_size//1024} KB")
print({k: len(v) for k, v in anom.items()})
